"""Text and metadata extraction for supported file types.

Extraction is mechanical and factual. Failure or an unsupported type never
invalidates or removes the preserved source; the outcome is recorded and the
file remains in the system with its Proof Object.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime

from . import config

STATUS_EXTRACTED = "extracted"
STATUS_UNSUPPORTED = "unsupported"
STATUS_FAILED = "failed"
STATUS_REGISTERED = "registered"  # images and archives: preserved + indexed by name/metadata only


@dataclass
class ExtractionResult:
    status: str
    text: str = ""
    error: str | None = None
    doc_type: str | None = None
    detected_people: str | None = None
    detected_date: str | None = None
    metadata: dict = field(default_factory=dict)


def _truncate(text: str) -> str:
    return text[: config.MAX_EXTRACTED_CHARS]


def _extract_pdf(data: bytes) -> ExtractionResult:
    """PDF text via pypdf; falls back to poppler's pdftotext when pypdf
    is not installed. Both are real extractors — there is no mock path."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return _extract_pdf_poppler(data)

    reader = PdfReader(io.BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    meta = {}
    if reader.metadata:
        meta = {k.lstrip("/"): str(v) for k, v in reader.metadata.items() if v}
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate("\n".join(parts)),
        doc_type="PDF document",
        metadata={"pages": len(reader.pages), "extractor": "pypdf", **meta},
    )


def _extract_pdf_poppler(data: bytes) -> ExtractionResult:
    import shutil
    import subprocess

    if not shutil.which("pdftotext"):
        raise RuntimeError(
            "no PDF extractor available: pypdf is not installed and pdftotext was not found"
        )
    proc = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", "-", "-"],
        input=data, capture_output=True, timeout=120, check=True,
    )
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate(proc.stdout.decode("utf-8", errors="replace")),
        doc_type="PDF document",
        metadata={"extractor": "pdftotext"},
    )


_DOCX_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _extract_docx(data: bytes) -> ExtractionResult:
    try:
        import docx
    except ImportError:
        return _extract_docx_stdlib(data)

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    props = document.core_properties
    meta = {
        k: str(getattr(props, k))
        for k in ("author", "created", "modified", "title")
        if getattr(props, k, None)
    }
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate("\n".join(parts)),
        doc_type="Word document",
        metadata={"extractor": "python-docx", **meta},
    )


def _extract_docx_stdlib(data: bytes) -> ExtractionResult:
    """DOCX is ZIP + XML; read word/document.xml paragraph text directly."""
    import xml.etree.ElementTree as ET
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    paragraphs = []
    for para in root.iter(f"{_DOCX_NS}p"):
        runs = [t.text or "" for t in para.iter(f"{_DOCX_NS}t")]
        if runs:
            paragraphs.append("".join(runs))
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate("\n".join(paragraphs)),
        doc_type="Word document",
        metadata={"extractor": "stdlib-ooxml"},
    )


def _extract_xlsx(data: bytes) -> ExtractionResult:
    try:
        import openpyxl
    except ImportError:
        return _extract_xlsx_stdlib(data)

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"[sheet] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append("\t".join(cells))
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate("\n".join(parts)),
        doc_type="Spreadsheet",
        metadata={"sheets": len(workbook.worksheets), "extractor": "openpyxl"},
    )


def _extract_xlsx_stdlib(data: bytes) -> ExtractionResult:
    """XLSX is ZIP + XML; pull shared strings and inline cell values."""
    import re as _re
    import xml.etree.ElementTree as ET
    import zipfile

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    parts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            sroot = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in sroot.iter(f"{ns}si"):
                shared.append("".join(t.text or "" for t in si.iter(f"{ns}t")))
        sheets = sorted(
            n for n in zf.namelist()
            if _re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)
        )
        for name in sheets:
            parts.append(f"[sheet] {name.rsplit('/', 1)[-1]}")
            root = ET.fromstring(zf.read(name))
            for row in root.iter(f"{ns}row"):
                cells = []
                for cell in row.iter(f"{ns}c"):
                    value = cell.find(f"{ns}v")
                    if value is None or value.text is None:
                        continue
                    if cell.get("t") == "s":
                        idx = int(value.text)
                        cells.append(shared[idx] if idx < len(shared) else "")
                    else:
                        cells.append(value.text)
                if cells:
                    parts.append("\t".join(cells))
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate("\n".join(parts)),
        doc_type="Spreadsheet",
        metadata={"sheets": len(sheets), "extractor": "stdlib-ooxml"},
    )


def _extract_csv(data: bytes) -> ExtractionResult:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    flat = "\n".join("\t".join(row) for row in rows)
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate(flat),
        doc_type="CSV table",
        metadata={"rows": len(rows)},
    )


def _extract_txt(data: bytes) -> ExtractionResult:
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate(data.decode("utf-8", errors="replace")),
        doc_type="Text file",
    )


def _extract_eml(data: bytes) -> ExtractionResult:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    headers = {k: str(msg.get(k, "")) for k in ("From", "To", "Cc", "Subject", "Date")}
    people = []
    for name, addr in getaddresses(
        [headers.get("From", ""), headers.get("To", ""), headers.get("Cc", "")]
    ):
        label = name or addr
        if label and label not in people:
            people.append(label)
        if name and addr and addr not in people:
            people.append(addr)
    detected_date = None
    if headers.get("Date"):
        try:
            detected_date = parsedate_to_datetime(headers["Date"]).date().isoformat()
        except (TypeError, ValueError):
            detected_date = None
    body = msg.get_body(preferencelist=("plain", "html"))
    body_text = body.get_content() if body else ""
    attachments = [
        part.get_filename() for part in msg.iter_attachments() if part.get_filename()
    ]
    text = "\n".join(
        [f"{k}: {v}" for k, v in headers.items() if v] + ["", str(body_text)]
    )
    return ExtractionResult(
        status=STATUS_EXTRACTED,
        text=_truncate(text),
        doc_type="Email message",
        detected_people="; ".join(people) or None,
        detected_date=detected_date,
        metadata={"headers": headers, "attachments": attachments},
    )


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".csv": _extract_csv,
    ".txt": _extract_txt,
    ".eml": _extract_eml,
}


def extract(data: bytes, ext: str) -> ExtractionResult:
    ext = ext.lower()
    if ext in config.IMAGE_EXTENSIONS:
        return ExtractionResult(
            status=STATUS_REGISTERED,
            doc_type="Image",
            metadata={"note": "image registered; no text extraction in v1.0"},
        )
    if ext in config.ARCHIVE_EXTENSIONS:
        return ExtractionResult(status=STATUS_REGISTERED, doc_type="ZIP archive")
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        return ExtractionResult(
            status=STATUS_UNSUPPORTED,
            error=f"unsupported file type '{ext or 'no extension'}': preserved without text extraction",
        )
    try:
        return extractor(data)
    except Exception as exc:  # the source is already preserved; record and continue
        return ExtractionResult(status=STATUS_FAILED, error=f"{type(exc).__name__}: {exc}")


def metadata_to_json(result: ExtractionResult) -> str:
    return json.dumps(result.metadata, default=str, ensure_ascii=False)
