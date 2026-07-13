"""Test fixtures for the WO-001 acceptance suite.

All fixture files are SYNTHETIC, generated at test runtime into temporary
directories. Nothing here is real case evidence, nothing is committed to
the repository, and nothing touches the live data root: every test runs
against an isolated EIW_DATA_DIR.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

try:
    import pytest
except ImportError:  # the standalone runner works without pytest
    pytest = None


def make_pdf(text: str) -> bytes:
    """Hand-assemble a minimal single-page PDF whose page text is extractable."""
    safe = text.replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(text: str) -> bytes:
    try:
        import docx
    except ImportError:
        return _make_docx_stdlib(text)

    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph(text)
    d.save(buf)
    return buf.getvalue()


def _make_docx_stdlib(text: str) -> bytes:
    """Assemble a minimal valid DOCX (OOXML zip) with one paragraph."""
    from xml.sax.saxutils import escape

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p></w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    return buf.getvalue()


def make_xlsx(rows: list[list[str]]) -> bytes:
    try:
        import openpyxl
    except ImportError:
        return _make_xlsx_stdlib(rows)

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(buf)
    return buf.getvalue()


def _make_xlsx_stdlib(rows: list[list[str]]) -> bytes:
    """Assemble a minimal valid XLSX (OOXML zip) using shared strings."""
    from xml.sax.saxutils import escape

    strings: list[str] = []
    row_xml = []
    for r, row in enumerate(rows, start=1):
        cells = []
        for c, value in enumerate(row):
            ref = chr(ord("A") + c) + str(r)
            strings.append(str(value))
            cells.append(f'<c r="{ref}" t="s"><v>{len(strings) - 1}</v></c>')
        row_xml.append(f'<row r="{r}">' + "".join(cells) + "</row>")
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<sheetData>" + "".join(row_xml) + "</sheetData></worksheet>"
    )
    sst = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        f' count="{len(strings)}" uniqueCount="{len(strings)}">'
        + "".join(f"<si><t>{escape(s)}</t></si>" for s in strings)
        + "</sst>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        "</Relationships>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("xl/workbook.xml", workbook)
        zf.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet)
        zf.writestr("xl/sharedStrings.xml", sst)
    return buf.getvalue()


def make_eml(sender: str, to: str, subject: str, date: str, body: str) -> bytes:
    return (
        f"From: {sender}\r\nTo: {to}\r\nSubject: {subject}\r\nDate: {date}\r\n"
        f"MIME-Version: 1.0\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n{body}\r\n"
    ).encode()


def make_png(shade: int = 0) -> bytes:
    """Build a valid 1x1 grayscale PNG (images are registered, not extracted)."""
    import struct
    import zlib

    def chunk(kind: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload)) + kind + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)
    idat = zlib.compress(bytes([0, shade & 0xFF]))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")
    )


# Registration-only formats: valid magic numbers are sufficient for v1.0.
JPEG_MIN = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00" + b"SYNTHETIC-TEST-DATA jpeg payload" + b"\xff\xd9"
TIFF_MIN = b"II*\x00\x08\x00\x00\x00" + b"SYNTHETIC-TEST-DATA tiff payload"


def synthetic_corpus() -> dict[str, bytes]:
    """26 mixed supported files + 1 unsupported + material for the ZIP test.

    Every document is watermarked SYNTHETIC-TEST-DATA so it can never be
    mistaken for real case evidence.
    """
    corpus: dict[str, bytes] = {}
    people = ["Jordan Example", "Casey Sample", "Riley Fixture", "Avery Mockman"]
    for i in range(1, 7):
        corpus[f"synthetic-letter-{i:02d}.pdf"] = make_pdf(
            f"SYNTHETIC-TEST-DATA letter {i} regarding the aquamarine ledger."
            f" Prepared by {people[i % 4]}."
        )
    for i in range(1, 6):
        corpus[f"synthetic-memo-{i:02d}.docx"] = make_docx(
            f"SYNTHETIC-TEST-DATA memo {i}: the turquoise settlement schedule"
            f" was reviewed by {people[(i + 1) % 4]}."
        )
    for i in range(1, 5):
        corpus[f"synthetic-note-{i:02d}.txt"] = (
            f"SYNTHETIC-TEST-DATA note {i}. The vermilion inventory belongs to"
            f" {people[(i + 2) % 4]}.\n"
        ).encode()
    corpus["synthetic-ledger-01.csv"] = (
        "date,item,amount\n2023-04-01,SYNTHETIC-TEST-DATA umbrella,12.50\n"
        "2023-04-02,periwinkle folder,3.25\n"
    ).encode()
    corpus["synthetic-ledger-02.csv"] = (
        "date,item,amount\n2023-05-01,SYNTHETIC-TEST-DATA lantern,99.00\n"
    ).encode()
    corpus["synthetic-accounts-01.xlsx"] = make_xlsx(
        [["quarter", "balance"], ["Q1", "1000"], ["Q2", "SYNTHETIC-TEST-DATA 2000"]]
    )
    corpus["synthetic-accounts-02.xlsx"] = make_xlsx(
        [["account", "note"], ["checking", "cerulean transfer SYNTHETIC-TEST-DATA"]]
    )
    for i, date in enumerate(
        ["Mon, 03 Apr 2023 10:15:00 -0400", "Tue, 06 Jun 2023 09:00:00 -0400",
         "Wed, 09 Aug 2023 16:45:00 -0400"], start=1,
    ):
        corpus[f"synthetic-mail-{i:02d}.eml"] = make_eml(
            f"{people[i % 4]} <sender{i}@example.test>",
            "Jordan Example <jordan@example.test>",
            f"SYNTHETIC-TEST-DATA thread {i} about the mulberry appraisal",
            date,
            f"Body {i}: the mulberry appraisal draft is attached in spirit. SYNTHETIC-TEST-DATA.",
        )
    corpus["synthetic-photo-01.png"] = make_png(0)
    corpus["synthetic-photo-02.jpg"] = JPEG_MIN
    corpus["synthetic-scan-01.tiff"] = TIFF_MIN
    corpus["synthetic-photo-03.png"] = make_png(255)  # distinct bytes
    # Unsupported type: must be preserved, registered, and reported — never lost.
    corpus["synthetic-unknown-01.xyz"] = b"SYNTHETIC-TEST-DATA opaque payload"
    return corpus


def make_zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


if pytest is not None:

    @pytest.fixture()
    def workstation(tmp_path: Path, monkeypatch):
        """A live-mode app instance bound to an isolated temporary data root."""
        from fastapi.testclient import TestClient

        data_dir = tmp_path / "eiw-data"
        monkeypatch.setenv("EIW_DATA_DIR", str(data_dir))
        monkeypatch.delenv("EIW_MODE", raising=False)

        from backend.app.main import create_app

        client = TestClient(create_app())
        yield client, data_dir
        client.app.state.conn.close()
